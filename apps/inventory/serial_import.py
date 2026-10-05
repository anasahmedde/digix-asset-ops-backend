"""Reading a column of serial numbers out of a spreadsheet.

Typing two hundred serials by hand is not a job anybody should be given,
so the store can hand over the list the supplier sent instead. The file is
read here rather than in the browser: the server already writes xlsx, so
it already has the library, and one reader means one set of rules about
what counts as a serial.

Nothing is saved from this — the serials go back to the screen that asked,
where they fill in the same boxes somebody would otherwise have typed, and
can still be corrected before anything is opened.
"""
import csv
import io

#: Headers that mean "this is the serial column". Matched case-insensitively
#: against the first row; anything else and the first column is used.
HEADERS = {"serial", "serial no", "serial no.", "serial number", "serial_number", "sn", "s/n"}

#: One upload cannot open more stock than the form allows anyway.
LIMIT = 500


def _clean(value) -> str:
    if value is None:
        return ""
    # A serial typed into Excel as a number comes back as a float: 12345.0
    # is not a serial number anybody printed on a label.
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _rows_from_csv(data: bytes):
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("That file is not text this can read.")
    return [row for row in csv.reader(io.StringIO(text))]


def _rows_from_excel(data: bytes):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = book.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def read_serials(upload):
    """Pull the serial numbers out of an uploaded sheet.

    Returns ``(serials, notes)`` — the list in file order with blanks and
    repeats dropped, and a few lines saying what was left out, so nobody
    has to compare the result against the file row by row to find out.
    """
    name = (getattr(upload, "name", "") or "").lower()
    data = upload.read()
    if not data:
        raise ValueError("That file is empty.")

    if name.endswith(".csv") or name.endswith(".txt"):
        rows = _rows_from_csv(data)
    elif name.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_excel(data)
    elif name.endswith(".xls"):
        raise ValueError(
            "That is the old .xls format. Save it as .xlsx or .csv and try again."
        )
    else:
        raise ValueError("Upload a .xlsx or .csv file.")

    if not rows:
        raise ValueError("There is nothing in that file.")

    # Which column holds them, and whether the first row is a heading.
    column = 0
    first = [_clean(c).lower() for c in rows[0]]
    heading = any(c in HEADERS for c in first)
    if heading:
        column = next(i for i, c in enumerate(first) if c in HEADERS)
        rows = rows[1:]

    serials, seen, blank, repeated = [], set(), 0, 0
    for row in rows:
        value = _clean(row[column]) if len(row) > column else ""
        if not value:
            blank += 1
            continue
        if value.lower() in seen:
            repeated += 1
            continue
        seen.add(value.lower())
        serials.append(value)

    if not serials:
        raise ValueError(
            "No serial numbers found. Put one per row in the first column, "
            "or head the column 'Serial number'."
        )

    notes = []
    if heading:
        notes.append("Read the column headed with a serial heading; the first row was its title.")
    if repeated:
        notes.append(f"{repeated} repeated serial{'s' if repeated > 1 else ''} dropped.")
    if blank:
        notes.append(f"{blank} blank row{'s' if blank > 1 else ''} skipped.")
    over = len(serials) - LIMIT
    if over > 0:
        serials = serials[:LIMIT]
        notes.append(f"Only the first {LIMIT} were taken; {over} more are in the file.")
    return serials, notes
