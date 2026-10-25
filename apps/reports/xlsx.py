"""Excel workbooks for reports, laid out the way finance and operations teams
expect to receive them: a title block saying what the sheet is and when it was
drawn, a frozen and filterable header, typed columns (money, quantities,
dates, percentages) with proper number formats, a totals row that is a live
SUM rather than a pasted figure, and A4-landscape print setup.

Columns are ``{"key", "label", "kind", "total"}``. ``kind`` is one of text,
int, qty (a count or measure, shown as typed), number, money, variance (money
where over plan is red and under is green), percent (0-100), date or datetime;
``total`` asks for a SUM in the totals row (money and variance columns are
totalled unless told otherwise).
"""

from __future__ import annotations

import datetime
from decimal import Decimal, InvalidOperation

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common.exports import XLSX_CONTENT_TYPE, _cell

INK = "1F2A37"
MUTED = "6B7280"
HEAD_FILL = "134E4A"
SECTION_FILL = "DCEBE9"
ZEBRA_FILL = "F4F8F8"
TOTAL_FILL = "E3EFED"
RULE = "D5DFDD"

FORMATS = {
    "money": "#,##0.00;[Red]-#,##0.00",
    "variance": "[Red]+#,##0.00;[Color10]-#,##0.00;0.00",
    "qty": "General",
    "number": "#,##0.00",
    "int": "#,##0",
    "percent": "0.0%",
    "date": "dd-mmm-yyyy",
    "datetime": "dd-mmm-yyyy hh:mm",
}
NUMERIC = {"money", "variance", "number", "int", "qty", "percent"}

_thin = Side(style="thin", color=RULE)
_head_font = Font(bold=True, color="FFFFFF", size=10)
_head_fill = PatternFill("solid", fgColor=HEAD_FILL)
_zebra = PatternFill("solid", fgColor=ZEBRA_FILL)
_total_fill = PatternFill("solid", fgColor=TOTAL_FILL)
_section_fill = PatternFill("solid", fgColor=SECTION_FILL)


def _value(value, kind):
    """A value as Excel should hold it for its column's kind."""
    value = _cell(value)
    if value == "":
        return None
    if kind in ("money", "variance", "number", "percent", "qty"):
        if isinstance(value, str):
            try:
                value = Decimal(value.replace(",", ""))
            except InvalidOperation:
                return value
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            value = float(value)
            if kind == "qty" and value.is_integer():
                return int(value)
            return value / 100 if kind == "percent" else value
    if kind == "date" and isinstance(value, datetime.datetime):
        return value.date()
    return value


def _width(value, kind):
    if value is None:
        return 0
    if kind in ("money", "variance"):
        return len(f"{value:,.2f}") + 2 if isinstance(value, float) else len(str(value))
    if kind in ("date",):
        return 11
    if kind == "datetime":
        return 17
    return len(str(value))


class ReportBook:
    """One workbook, one sheet per table or summary."""

    def __init__(self, title: str, user=None):
        self.wb = Workbook()
        self._fresh = True
        self.title = title
        who = (user.get_full_name() or user.username) if user is not None and user.is_authenticated else ""
        self.drawn = f"Generated {timezone.localtime():%d %b %Y, %H:%M}" + (f" by {who}" if who else "")
        self.wb.properties.title = title
        self.wb.properties.creator = who or "Asset Ops"

    def _sheet(self, name: str):
        if self._fresh:
            ws = self.wb.active
            ws.title = name[:31]
            self._fresh = False
        else:
            ws = self.wb.create_sheet(name[:31])
        ws.sheet_view.showGridLines = False
        return ws

    def _heading(self, ws, title, subtitle) -> int:
        ws.cell(1, 1, title).font = Font(bold=True, size=14, color=INK)
        row = 2
        for line in subtitle:
            if line:
                ws.cell(row, 1, line).font = Font(size=9, color=MUTED)
                row += 1
        ws.cell(row, 1, self.drawn).font = Font(size=9, italic=True, color=MUTED)
        return row + 2

    @staticmethod
    def _print_setup(ws, header_row=None, columns=0):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        # Wide registers print across pages rather than shrinking out of reach.
        ws.page_setup.fitToWidth = max(1, -(-columns // 12))
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddFooter.left.text = "&A"
        ws.oddFooter.right.text = "Page &P of &N"
        if header_row:
            ws.print_title_rows = f"{header_row}:{header_row}"

    def table(self, name, title, columns, rows, subtitle=(), totals=True, empty="Nothing to report."):
        ws = self._sheet(name)
        head = self._heading(ws, title, subtitle)
        widths = [len(c["label"]) + 2 for c in columns]
        for i, col in enumerate(columns, start=1):
            cell = ws.cell(head, i, col["label"])
            cell.font = _head_font
            cell.fill = _head_fill
            cell.border = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)
            cell.alignment = Alignment(
                horizontal="right" if col.get("kind") in NUMERIC else "left", vertical="center", wrap_text=True,
            )
        ws.row_dimensions[head].height = 30

        r = head
        for n, row in enumerate(rows):
            r = head + 1 + n
            for i, col in enumerate(columns, start=1):
                kind = col.get("kind", "text")
                value = _value(row.get(col["key"]), kind)
                cell = ws.cell(r, i, value)
                cell.border = Border(bottom=_thin)
                cell.alignment = Alignment(vertical="top", wrap_text=kind == "text")
                if kind in FORMATS:
                    cell.number_format = FORMATS[kind]
                if n % 2:
                    cell.fill = _zebra
                widths[i - 1] = max(widths[i - 1], _width(value, kind) + 2)
        last = r
        if not rows:
            ws.cell(head + 1, 1, empty).font = Font(italic=True, color=MUTED)
            last = head + 1

        summed = [i for i, c in enumerate(columns, start=1) if c.get("total", c.get("kind") in ("money", "variance"))]
        if totals and rows and summed:
            t = last + 1
            for i in range(1, len(columns) + 1):
                cell = ws.cell(t, i)
                cell.fill = _total_fill
                cell.font = Font(bold=True, color=INK)
                cell.border = Border(top=Side(style="thin", color=INK), bottom=Side(style="double", color=INK))
            ws.cell(t, 1, "Total")
            for i in summed:
                letter = get_column_letter(i)
                cell = ws.cell(t, i, f"=SUBTOTAL(9,{letter}{head + 1}:{letter}{last})")
                cell.number_format = FORMATS.get(columns[i - 1].get("kind"), "#,##0.00")

        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(8, min(60, w))
        ws.freeze_panes = ws.cell(head + 1, 2 if len(columns) > 3 else 1)
        if rows:
            ws.auto_filter.ref = f"A{head}:{get_column_letter(len(columns))}{last}"
        self._print_setup(ws, head, len(columns))
        return ws

    def summary(self, name, title, sections, subtitle=()):
        """Label / value pairs under section bands: ``[(section, [(label, value, kind)])]``."""
        ws = self._sheet(name)
        r = self._heading(ws, title, subtitle)
        for section, pairs in sections:
            band = ws.cell(r, 1, section)
            band.font = Font(bold=True, color=INK)
            for col in (1, 2):
                ws.cell(r, col).fill = _section_fill
            r += 1
            for label, value, kind in pairs:
                ws.cell(r, 1, label).font = Font(color=MUTED)
                cell = ws.cell(r, 2, _value(value, kind))
                cell.alignment = Alignment(horizontal="left", wrap_text=True)
                if kind in FORMATS:
                    cell.number_format = FORMATS[kind]
                ws.cell(r, 1).border = ws.cell(r, 2).border = Border(bottom=_thin)
                r += 1
            r += 1
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 52
        self._print_setup(ws)
        ws.page_setup.orientation = "portrait"
        return ws

    def response(self, filename: str) -> HttpResponse:
        dated = f"{filename}-{timezone.localdate().isoformat()}.xlsx"
        resp = HttpResponse(content_type=XLSX_CONTENT_TYPE)
        resp["Content-Disposition"] = f'attachment; filename="{dated}"'
        self.wb.save(resp)
        return resp
