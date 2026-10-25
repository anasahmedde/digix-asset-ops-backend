import pytest
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.assets.models import AssetType, Device
from apps.maintenance.models import MaintenanceSchedule, MaintenanceVisit
from apps.sites.models import Site


@pytest.mark.django_db
def test_visits_report_lists_each_visit_and_rolls_up_by_technician():
    ops = User.objects.create_user(username="rep-ops", password="x", role="ops_manager")
    tech = User.objects.create_user(username="rep-tech", password="x", role="technician",
                                    first_name="Rana", last_name="Tech")
    site = Site.objects.create(name="Report Site", address="x")
    kind = AssetType.objects.create(name="Rep Screen")
    device = Device.objects.create(asset_type=kind, asset_code="AST-REP-1", status=Device.Status.ACTIVE)
    job = MaintenanceSchedule.objects.create(
        title="Monthly clean", device=device, site=site, assigned_to=tech,
        next_due=timezone.localdate(), maintenance_type="preventive",
    )
    now = timezone.now()
    for i in range(2):
        MaintenanceVisit.objects.create(
            schedule=job, due_date=timezone.localdate(), assigned_to=tech, sequence=i + 1,
            status=MaintenanceVisit.Status.COMPLETED, started_at=now, completed_at=now + timezone.timedelta(hours=2),
            asset_verified_at=now if i == 0 else None,
        )
    c = APIClient()
    c.force_authenticate(ops)

    r = c.get("/api/reports/generate/", {"type": "visits"})
    assert r.status_code == 200, r.content
    keys = [col["key"] for col in r.data["columns"]]
    assert keys[:3] == ["technician", "site", "project"]
    assert len(r.data["rows"]) >= 2
    assert {row["asset_verified"] for row in r.data["rows"] if row["asset"] == "AST-REP-1"} == {"Yes", "No"}

    r = c.get("/api/reports/generate/", {"type": "visits", "group_by": "technician", "technician": str(tech.id)})
    assert r.status_code == 200, r.content
    # The schedule opened its own planned round as well as the two completed ones.
    assert r.data["rows"] == [{
        "technician": "Rana Tech", "visits": 3, "completed": 2, "in_progress": 0,
        "awaiting_review": 0, "planned": 1, "unresolved": 0, "hours": 4.0, "avg_hours": 2.0,
    }]

    r = c.get("/api/reports/generate/", {"type": "visits", "group_by": "site", "site": str(site.id)})
    assert r.data["rows"][0]["site"] == "Report Site" and r.data["rows"][0]["visits"] == 3


@pytest.mark.django_db
def test_the_catalog_and_every_report_follow_the_rights():
    """Each report is offered to, and runs for, only those with its right; the
    same report comes back as a formatted workbook; a stranger is refused."""
    from io import BytesIO

    from openpyxl import load_workbook

    from apps.reports.builders import REPORTS

    admin = User.objects.create_superuser(username="rep-admin", password="x", email="rep-admin@x.test")
    tech = User.objects.create_user(username="rep-tech2", password="x", role="technician")
    c = APIClient()
    c.force_authenticate(admin)
    keys = [r["key"] for r in c.get("/api/reports/catalog/").data["reports"]]
    assert set(keys) == set(REPORTS)
    for key in keys:
        r = c.get("/api/reports/generate/", {"type": key})
        assert r.status_code == 200, (key, r.content)
        assert r.data["columns"] and "kpis" in r.data
        x = c.get("/api/reports/generate/", {"type": key, "export": "xlsx"})
        assert x.status_code == 200, key
        ws = load_workbook(BytesIO(x.content)).active
        assert ws["A1"].value == REPORTS[key]["title"]

    t = APIClient()
    t.force_authenticate(tech)
    offered = {r["key"] for r in t.get("/api/reports/catalog/").data["reports"]}
    assert "purchase_orders" not in offered
    assert t.get("/api/reports/generate/", {"type": "purchase_orders"}).status_code == 403


@pytest.mark.django_db
def test_the_project_workbook_has_every_sheet_and_hides_costs_without_the_right():
    from io import BytesIO

    from openpyxl import load_workbook

    from apps.teams.models import Project

    admin = User.objects.create_superuser(username="rep-admin2", password="x", email="rep-admin2@x.test")
    project = Project.objects.create(name="Workbook Project")
    Device.objects.create(asset_type=AssetType.objects.create(name="WB Screen"), project=project, asset_code="AST-WB-1")
    c = APIClient()
    c.force_authenticate(admin)
    r = c.get(f"/api/reports/project/{project.pk}/")
    assert r.status_code == 200, r.content
    assert project.project_code in r["Content-Disposition"]
    sheets = load_workbook(BytesIO(r.content)).sheetnames
    assert sheets == ["Summary", "Cost", "Components", "Production", "Assets & installation",
                      "Site visits", "Overheads", "Milestones", "Activity"]
    assets = load_workbook(BytesIO(r.content))["Assets & installation"]
    assert any(cell.value == "AST-WB-1" for row in assets.iter_rows() for cell in row)

    # Marketing sees projects but not what they cost.
    viewer = User.objects.create_user(username="rep-viewer", password="x", role="marketing")
    v = APIClient()
    v.force_authenticate(viewer)
    r = v.get(f"/api/reports/project/{project.pk}/")
    assert r.status_code == 200, r.content
    assert load_workbook(BytesIO(r.content)).sheetnames == [
        "Summary", "Assets & installation", "Site visits", "Milestones", "Activity",
    ]
